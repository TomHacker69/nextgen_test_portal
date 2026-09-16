from fastapi import APIRouter, Depends, Query
from typing import Optional
from bson import ObjectId
from fastapi.responses import StreamingResponse
import csv
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from app.core.deps import require_admin
from app.core.database import get_db

router = APIRouter(prefix="/api/admin/export", tags=["admin-export"])


@router.get("/{exam_id}/leaderboard/csv")
async def export_leaderboard_csv(exam_id: str, admin=Depends(require_admin), db=Depends(get_db)):
    """Export full leaderboard as CSV."""
    rows = await _get_leaderboard_rows(exam_id, db)

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Rank", "Name", "RollNumber", "Email", "Score", "MaxScore", "Accuracy", "Correct", "Wrong", "Unattempted", "TimeTaken(s)", "Status", "SubmittedAt"])

    for r in rows:
        writer.writerow([
            r["rank"], r["name"], r.get("roll_number", ""), r.get("email", ""),
            r["score"], r["max_score"], f"{r['accuracy']}%",
            r["correct"], r["wrong"], r["unattempted"],
            r.get("time_taken_seconds", ""), r["status"],
            r.get("submitted_at", ""),
        ])

    output.seek(0)
    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename=leaderboard_{exam_id}.csv"},
    )


@router.get("/{exam_id}/leaderboard/xlsx")
async def export_leaderboard_xlsx(exam_id: str, admin=Depends(require_admin), db=Depends(get_db)):
    """Export full leaderboard as XLSX with multiple sheets."""
    exam = await db.exams.find_one({"_id": ObjectId(exam_id)})
    rows = await _get_leaderboard_rows(exam_id, db)

    wb = openpyxl.Workbook()

    # Sheet 1: Leaderboard
    ws = wb.active
    ws.title = "Leaderboard"
    headers = ["Rank", "Name", "Roll Number", "Email", "Score", "Max Score", "Accuracy %", "Correct", "Wrong", "Unattempted", "Time (s)", "Status", "Submitted At"]
    _write_header_row(ws, headers)

    for r in rows:
        ws.append([
            r["rank"], r["name"], r.get("roll_number", ""), r.get("email", ""),
            r["score"], r["max_score"], r["accuracy"],
            r["correct"], r["wrong"], r["unattempted"],
            r.get("time_taken_seconds"), r["status"],
            str(r.get("submitted_at", "")),
        ])

    # Sheet 2: Exam Info
    ws2 = wb.create_sheet("Exam Info")
    ws2.append(["Field", "Value"])
    ws2.append(["Exam ID", exam_id])
    ws2.append(["Title", exam.get("title", "") if exam else ""])
    ws2.append(["Total Candidates", len(rows)])
    ws2.append(["Submitted", sum(1 for r in rows if r["status"] == "submitted")])

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        iter([output.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=leaderboard_{exam_id}.xlsx"},
    )


def _write_header_row(ws, headers):
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="1E293B")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")


async def _get_leaderboard_rows(exam_id: str, db):
    pipeline = [
        {"$match": {"exam_id": exam_id, "status": {"$in": ["submitted", "expired"]}}},
        {"$sort": {"score": -1, "correct_count": -1, "time_taken_seconds": 1, "_id": 1}},
        {"$lookup": {
            "from": "users",
            "let": {"sid": {"$toObjectId": "$student_id"}},
            "pipeline": [{"$match": {"$expr": {"$eq": ["$_id", "$$sid"]}}}],
            "as": "student",
        }},
        {"$unwind": {"path": "$student", "preserveNullAndEmpty": True}},
    ]

    attempts = await db.attempts.aggregate(pipeline).to_list(None)
    rows = []
    for i, a in enumerate(attempts):
        student = a.get("student", {})
        correct = a.get("correct_count", 0)
        wrong = a.get("wrong_count", 0)
        accuracy = round(correct / max(correct + wrong, 1) * 100, 1) if (correct + wrong) > 0 else 0

        rows.append({
            "rank": i + 1,
            "name": student.get("name", "Unknown"),
            "roll_number": student.get("roll_number"),
            "email": student.get("email"),
            "score": a.get("score", 0),
            "max_score": a.get("max_score", 0),
            "accuracy": accuracy,
            "correct": correct,
            "wrong": wrong,
            "unattempted": a.get("unattempted_count", 0),
            "time_taken_seconds": a.get("time_taken_seconds"),
            "status": a.get("status"),
            "submitted_at": a.get("submitted_at"),
            "attempt_id": str(a["_id"]),
        })

    return rows
